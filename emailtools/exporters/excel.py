from __future__ import annotations
import logging
import math
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from emailtools.extractors.tables import (normalize_value, MAX_EXCEL_CELL_CHARS, ILLEGAL_XML_CHARS, WORKSHEET_NAME)

def safe_excel_value(
    value: object,
    logger: logging.Logger,
    context: str,
) -> str:
    text = normalize_value(value)
    cleaned = ILLEGAL_XML_CHARS.sub(" ", text)
    if cleaned != text:
        logger.warning("Excel illegal control characters replaced: %s", context)
    if len(cleaned) > MAX_EXCEL_CELL_CHARS:
        logger.warning(
            "Excel cell truncated from %d to %d characters: %s",
            len(cleaned),
            MAX_EXCEL_CELL_CHARS,
            context,
        )
        return cleaned[: MAX_EXCEL_CELL_CHARS - 3] + "..."
    return cleaned


def _display_width(value: object) -> int:
    lines = str(value or "").splitlines() or [""]
    return max(sum(2 if ord(char) > 127 else 1 for char in line) for line in lines)


def _force_literal_text(cell) -> None:
    # openpyxl interprets strings beginning with '=' as formulas.  Every value
    # in this workbook comes from an untrusted email and must remain inert data.
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"


def write_excel(
    records: list[dict[str, str]],
    columns: list[str],
    output_path: Path,
    logger: logging.Logger,
    *, progress=None, is_cancelled=None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = WORKSHEET_NAME
    worksheet.sheet_view.showGridLines = False

    worksheet.append(columns)
    for row_number, record in enumerate(records, 2):
        if is_cancelled and is_cancelled():
            workbook.close()
            raise InterruptedError("Excel 저장을 중지했습니다.")
        worksheet.append(
            [
                safe_excel_value(
                    record.get(column, ""),
                    logger,
                    f"{record.get('source_eml', '')} / {column}",
                )
                for column in columns
            ]
        )
        if progress and (row_number % 50 == 0 or row_number == len(records) + 1):
            progress(row_number - 1, len(records))

    last_column = get_column_letter(len(columns))
    last_row = max(1, len(records) + 1)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{last_column}{last_row}"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        _force_literal_text(cell)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 28

    for row in worksheet.iter_rows(min_row=2, max_row=last_row):
        for cell in row:
            _force_literal_text(cell)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    preferred_widths = {
        "source_eml": 28,
        "_status": 24,
        "_error": 40,
        "subject": 32,
        "from": 26,
        "to": 26,
        "date": 24,
        "docx_filename": 28,
    }
    column_widths: dict[int, float] = {}
    for column_index, column_name in enumerate(columns, 1):
        measured = max(
            _display_width(worksheet.cell(row=row_index, column=column_index).value)
            for row_index in range(1, last_row + 1)
        )
        width = max(10, preferred_widths.get(column_name, 10), measured + 2)
        width = min(50, width)
        column_widths[column_index] = width
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(2, last_row + 1):
        if is_cancelled and is_cancelled():
            workbook.close()
            raise InterruptedError("Excel 저장을 중지했습니다.")
        estimated_lines = 1
        for column_index in range(1, len(columns) + 1):
            value = str(worksheet.cell(row=row_index, column=column_index).value or "")
            width = max(1, column_widths[column_index])
            wrapped_lines = sum(
                max(1, math.ceil(_display_width(line) / width))
                for line in (value.splitlines() or [""])
            )
            estimated_lines = max(estimated_lines, wrapped_lines)
        worksheet.row_dimensions[row_index].height = min(90, max(18, estimated_lines * 15))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=".tmp.xlsx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
        workbook.close()
