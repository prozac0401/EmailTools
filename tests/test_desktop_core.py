from __future__ import annotations
import copy
import logging
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from email.message import EmailMessage

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from emailtools import pipeline, reader
from emailtools.demo import create_demo
from emailtools.plans import options_for_mode, ExportPlan
from emailtools.models import ProcessingOptions
from emailtools.field_catalog import SourceHTMLParser
from emailtools.exporters.files import export_results
from emailtools.table_config import project_records, validate_columns
from openpyxl import load_workbook


class DesktopCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "메일"
        self.source.mkdir()
        self.logger = logging.Logger("desktop-test")

    def tearDown(self):
        self.temp.cleanup()

    def mail(self, name, html, *, subject="제목"):
        msg = EmailMessage()
        msg["Subject"] = subject
        msg["From"] = "sender@example.com"
        msg["To"] = "receiver@example.com"
        msg.set_content("본문")
        msg.add_alternative(html, subtype="html")
        path = self.source / name
        path.write_bytes(msg.as_bytes())
        return path

    def analyze(self, options=None, **kwargs):
        return pipeline.analyze(self.source, options or options_for_mode("analyze"), self.root / "spool", self.logger, **kwargs)

    def test_list_does_not_decode_render_or_extract_and_saves_excel(self):
        create_demo(self.source)
        with patch.object(reader, "attachment_bytes", side_effect=AssertionError("decoded")), \
             patch.object(reader, "render_mail_text", side_effect=AssertionError("rendered")), \
             patch("emailtools.extractors.tables.process_message_record", side_effect=AssertionError("tables")):
            result = self.analyze(options_for_mode("list"))
        self.assertEqual((result.success, result.failures), (3, 0))
        self.assertTrue(all(a.size is None for m in result.mails for a in m.attachments))
        folder = export_results(result, self.root / "out", self.logger)
        self.assertTrue((folder / "MailList.xlsx").exists())
        self.assertFalse((folder / "tables").exists())
        book = load_workbook(folder / "MailList.xlsx")
        self.assertEqual(book.active.max_row, 4)
        self.assertIn("to", [c.value for c in book.active[1]])
        book.close()

    def test_provenance_repetition_scopes_and_literal_pipe_values(self):
        table = "<table><tr><th>항목</th><th>값</th></tr><tr><td>공통</td><td>A | B</td></tr><tr><td>반복</td><td>동일</td></tr><tr><td>반복</td><td>동일</td></tr></table>"
        self.mail("1.eml", table)
        self.mail("2.eml", table, subject="다른 메일")
        result = self.analyze()
        by_name = {f["name"]: f for f in result.catalog.query()["fields"]}
        self.assertIn("common", by_name["공통"]["flags"])
        self.assertIn("repeated", by_name["반복"]["flags"])
        self.assertEqual(by_name["반복"]["occurrence_count"], 4)
        self.assertEqual(result.records[0]["공통"], "A | B")
        self.assertEqual(result.records[0]["반복"], "동일")

    def test_identical_eml_keeps_rows_but_cannot_manufacture_common_fields(self):
        first = self.mail("1.eml", "<table><tr><td>항목</td><td>값</td></tr><tr><td>공통</td><td>하나</td></tr></table>")
        (self.source / "2.eml").write_bytes(first.read_bytes())
        (self.source / "3.eml").write_bytes(first.read_bytes())
        result = self.analyze()
        field, = result.catalog.query()["fields"]
        self.assertEqual(len(result.mails), 3)
        self.assertEqual(field["mail_count"], 3)
        self.assertEqual(field["effective_mail_count"], 1)
        self.assertNotIn("common", field["flags"])

    def test_records_not_recommended_override_and_undo(self):
        table = "<table><tr><th>이름</th><th>부서</th></tr><tr><td>홍길동</td><td>개발</td></tr><tr><td>김민지</td><td>운영</td></tr></table>"
        self.mail("1.eml", table)
        self.mail("2.eml", table, subject="다른 메일")
        result = self.analyze()
        self.assertEqual(result.catalog.query()["counts"]["common"], 0)
        self.assertTrue(all(t.shape == "records" for t in result.catalog.tables))
        for table in result.catalog.tables:
            result.catalog.override_table(table.table_id, "key_value")
        self.assertGreater(result.catalog.query()["counts"]["common"], 0)
        for table in result.catalog.tables:
            result.catalog.override_table(table.table_id, "")
        self.assertEqual(result.catalog.query()["counts"]["common"], 0)

    def test_nested_and_merged_html_cells_owned_once(self):
        parser = SourceHTMLParser()
        parser.feed("<table><tr><td rowspan='2'>외부<table><tr><td>내부</td><td>값</td></tr></table></td><td>외부값</td></tr><tr><td>둘째</td></tr></table>")
        self.assertEqual(len(parser.tables), 2)
        outer = parser.tables[0]["cells"]
        self.assertNotIn("내부", outer[0].text)
        self.assertEqual(outer[-1].column, 1)
        self.assertEqual(sum(c.text == "내부" for t in parser.tables for c in t["cells"]), 1)

    def test_multiple_paths_read_once_and_table_only_ignores_other_payloads(self):
        create_demo(self.source)
        from emailtools.extractors.attachments import AttachmentSource
        original = reader.read_message
        with patch.object(reader, "read_message", wraps=original) as read:
            result = self.analyze(ProcessingOptions(mail_index=True), sources=[self.source, next(self.source.glob("*.eml"))])
        self.assertEqual(read.call_count, 3)
        self.assertEqual(result.attachment_count, 2)
        self.assertFalse(any(a.payload_path for m in result.mails for a in m.attachments))

    def test_cancellation_cleans_staging_and_preserves_previous_export(self):
        create_demo(self.source)
        result = self.analyze(options_for_mode("archive"))
        destination = self.root / "out"
        first = export_results(result, destination, self.logger, archive=True)
        cancelled = False
        def status(data):
            nonlocal cancelled
            if data["phase"] == "ZIP 생성":
                cancelled = True
        with self.assertRaises(InterruptedError):
            export_results(result, destination, self.logger, archive=True, progress=status, is_cancelled=lambda: cancelled)
        self.assertEqual(list(destination.iterdir()), [first])
        self.assertTrue((first / "MailList.xlsx").exists())

    def test_actual_phase_counts_and_export_capability_validation(self):
        create_demo(self.source)
        events = []
        result = self.analyze(options_for_mode("list"), status=events.append)
        finished = [e["completed"] for e in events if e["phase"] == "메일 처리"]
        self.assertEqual(finished, [1, 2, 3])
        self.assertIsNone(events[0]["total"])
        with self.assertRaises(ValueError):
            ExportPlan(attachments=True).apply(result)
        with self.assertRaises(InterruptedError):
            self.analyze(is_cancelled=lambda: True)

    def test_selection_projection_matches_export_with_custom_literal(self):
        create_demo(self.source)
        result = self.analyze()
        specs = [dict(source=None, name="검토", value="=1+1", enabled=True),
                 dict(source="교육명", name="교육", enabled=True),
                 dict(source="subject", name="메일 제목", enabled=True)]
        projected = project_records(result.records, validate_columns(specs, result.columns), self.logger)
        folder = export_results(result, self.root / "out", self.logger, columns=specs)
        book = load_workbook(folder / "tables/EML_Table_Result.xlsx")
        rows = list(book.active.values)
        self.assertEqual(list(rows[0]), ["검토", "교육", "메일 제목"])
        self.assertEqual(list(rows[1]), list(projected[0].values()))
        self.assertEqual(book.active["A2"].data_type, "s")
        book.close()


if __name__ == "__main__":
    unittest.main()
