from __future__ import annotations

import importlib.util
import subprocess
import sys
import tempfile
import unittest
import zipfile
from email.message import EmailMessage
from email.policy import SMTP
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape


REPO_ROOT = Path(__file__).resolve().parents[1]
APP_PATH = REPO_ROOT / "eml_table_to_excel" / "app.py"
RUN_BAT = REPO_ROOT / "eml_table_to_excel" / "run.bat"
EMBEDDED_PYTHON = REPO_ROOT / "eml_attachment_tool" / "python" / "python.exe"

SPEC = importlib.util.spec_from_file_location("eml_table_to_excel_app", APP_PATH)
assert SPEC and SPEC.loader
app = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = app
SPEC.loader.exec_module(app)

from openpyxl import load_workbook  # noqa: E402  (vendored path is set by app)


def docx_table_bytes(rows: list[list[str]]) -> bytes:
    table_rows: list[str] = []
    for row in rows:
        cells = "".join(
            "<w:tc><w:p><w:r><w:t>"
            + escape(cell)
            + "</w:t></w:r></w:p></w:tc>"
            for cell in row
        )
        table_rows.append(f"<w:tr>{cells}</w:tr>")
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:tbl>{''.join(table_rows)}</w:tbl></w:body></w:document>"
    )
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_xml.encode("utf-8"))
    return buffer.getvalue()


def html_table(rows: list[list[str]]) -> str:
    rendered_rows = []
    for row in rows:
        cells = "".join(f"<td>{escape(cell)}</td>" for cell in row)
        rendered_rows.append(f"<tr>{cells}</tr>")
    return "<html><body><table>" + "".join(rendered_rows) + "</table></body></html>"


def write_eml(
    path: Path,
    *,
    subject: str,
    email_rows: list[list[str]] | None,
    docx_rows: list[list[str]] | None,
    docx_filename: str = "평가 결과.docx",
    docx_payload: bytes | None = None,
) -> None:
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = "sender@example.com"
    message["To"] = "receiver@example.com"
    message["Date"] = "Thu, 03 Sep 2026 10:00:00 +0900"
    message.set_content("Synthetic fixture")
    if email_rows is not None:
        message.add_alternative(html_table(email_rows), subtype="html")
    if docx_rows is not None or docx_payload is not None:
        payload = docx_payload if docx_payload is not None else docx_table_bytes(docx_rows or [])
        message.add_attachment(
            payload,
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=docx_filename,
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(message.as_bytes(policy=SMTP))


class EmlTableToExcelTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="EmailTools_table_test_")
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_key_normalization_and_duplicate_values(self) -> None:
        self.assertEqual(app.normalize_key("\xa0 교육명 \r\n : \t"), "교육명")
        self.assertEqual(app.normalize_key("과정   평가"), "과정 평가")
        self.assertNotEqual(app.normalize_key("교육기간"), app.normalize_key("교육 기간"))
        values = app.table_to_key_value_dict(
            [
                ["항목명", "값"],
                ["강사", "김철수", "강사", "홍길동"],
                ["강사", "김철수"],
            ]
        )
        self.assertEqual(values, {"강사": "김철수 | 홍길동"})

    def test_html_two_and_four_cell_rows(self) -> None:
        tables = app.parse_html_tables(
            "<table><tr><th>항목명</th><th>값</th></tr>"
            "<tr><td>교육명 :</td><td>AI Basic</td><td>강사</td><td>김철수</td></tr>"
            "<tr><td>주요 의견</td><td>실습<br>유익</td></tr></table>"
        )
        self.assertEqual(len(tables), 1)
        self.assertEqual(
            app.table_to_key_value_dict(tables[0]),
            {"교육명": "AI Basic", "강사": "김철수", "주요 의견": "실습\n유익"},
        )

    def test_docx_table_is_parsed_from_bytes(self) -> None:
        payload = docx_table_bytes(
            [["과정 평가", "만족"], ["주요 의견", "실습이 좋음"]]
        )
        tables = app.parse_docx_tables(payload)
        self.assertEqual(len(tables), 1)
        self.assertEqual(
            app.table_to_key_value_dict(tables[0]),
            {"과정 평가": "만족", "주요 의견": "실습이 좋음"},
        )

    def test_dynamic_columns_order_mapping_missing_cells_and_collision(self) -> None:
        write_eml(
            self.root / "001.eml",
            subject="교육 A",
            email_rows=[
                [" 교육명 : ", "AI Basic"],
                ["교육기간", "2026.09.01"],
                ["강사", "김철수"],
            ],
            docx_rows=[
                ["과정 평가", "만족"],
                ["주요 의견", "실습이 좋음"],
                ["교육명", "AI Advanced"],
            ],
        )
        write_eml(
            self.root / "하위" / "002.eml",
            subject="교육 B",
            email_rows=[
                ["강사", "홍길동"],
                ["교육명", "PyTorch Advanced"],
                ["교육장", "B114"],
            ],
            docx_rows=[
                ["주요 의견", "실습 비중이 좋았음"],
                ["과정 평가", "만족"],
            ],
        )

        result = app.run_batch(self.root)
        expected_columns = app.SYSTEM_COLUMNS + [
            "교육명",
            "교육기간",
            "강사",
            "과정 평가",
            "주요 의견",
            "교육장",
        ]
        self.assertEqual(result.columns, expected_columns)
        self.assertEqual((result.total, result.success, result.warnings, result.failures), (2, 2, 0, 0))
        self.assertEqual(result.records[0]["교육명"], "AI Basic | AI Advanced")
        self.assertNotIn("교육기간", result.records[1])

        workbook = load_workbook(result.output_path)
        self.assertEqual(workbook.sheetnames, ["Result"])
        worksheet = workbook["Result"]
        headers = [cell.value for cell in worksheet[1]]
        self.assertEqual(headers, expected_columns)
        header_index = {name: index + 1 for index, name in enumerate(headers)}
        self.assertEqual(worksheet.cell(2, header_index["교육명"]).value, "AI Basic | AI Advanced")
        self.assertEqual(worksheet.cell(3, header_index["교육명"]).value, "PyTorch Advanced")
        self.assertEqual(worksheet.cell(3, header_index["교육장"]).value, "B114")
        self.assertIsNone(worksheet.cell(3, header_index["교육기간"]).value)
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(
            worksheet.auto_filter.ref,
            f"A1:{app.get_column_letter(worksheet.max_column)}3",
        )
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertTrue(worksheet["A2"].alignment.wrap_text)
        self.assertTrue(all(10 <= dimension.width <= 50 for dimension in worksheet.column_dimensions.values()))
        workbook.close()
        self.assertTrue(result.log_path.is_file())

    def test_warning_statuses_do_not_stop_batch(self) -> None:
        write_eml(
            self.root / "001_no_tables.eml",
            subject="No tables",
            email_rows=None,
            docx_rows=None,
        )
        write_eml(
            self.root / "002_bad_docx.eml",
            subject="Bad DOCX",
            email_rows=[["교육명", "Still parsed"]],
            docx_rows=None,
            docx_payload=b"not a zip file",
        )
        result = app.run_batch(self.root)
        self.assertEqual((result.total, result.success, result.warnings, result.failures), (2, 0, 2, 0))
        self.assertEqual(result.records[0]["_status"], "NO EMAIL TABLE | NO DOCX")
        self.assertEqual(result.records[1]["_status"], "DOCX PARSE ERROR")
        self.assertIn("BadZipFile", result.records[1]["_error"])
        log_text = result.log_path.read_text(encoding="utf-8-sig")
        self.assertIn("Traceback", log_text)
        self.assertIn("DOCX PARSE ERROR", log_text)
        self.assertTrue(result.output_path.is_file())

    def test_excel_cell_length_and_control_char_are_safe(self) -> None:
        logger = app.configure_logger(self.root / "length.log")
        try:
            output_path = self.root / "length.xlsx"
            record = {column: "" for column in app.SYSTEM_COLUMNS}
            record.update(
                {
                    "source_eml": "long.eml",
                    "_status": "OK",
                    "긴 값": "A" * 40000 + "\x01",
                    "formula_like": "=HYPERLINK(\"https://example.invalid\",\"click\")",
                }
            )
            columns = app.collect_columns([record])
            app.write_excel([record], columns, output_path, logger)
        finally:
            app.close_logger(logger)
        workbook = load_workbook(output_path)
        worksheet = workbook["Result"]
        column_index = [cell.value for cell in worksheet[1]].index("긴 값") + 1
        value = worksheet.cell(2, column_index).value
        self.assertEqual(len(value), app.MAX_EXCEL_CELL_CHARS)
        self.assertTrue(value.endswith("..."))
        self.assertNotIn("\x01", value)
        formula_column = [cell.value for cell in worksheet[1]].index("formula_like") + 1
        formula_cell = worksheet.cell(2, formula_column)
        self.assertEqual(formula_cell.data_type, "s")
        self.assertTrue(formula_cell.value.startswith("=HYPERLINK"))
        workbook.close()
        self.assertIn("truncated", (self.root / "length.log").read_text(encoding="utf-8-sig"))

    @unittest.skipUnless(EMBEDDED_PYTHON.is_file(), "embedded Python is not configured")
    def test_run_bat_uses_shared_embedded_python_from_other_cwd(self) -> None:
        input_root = self.root / "한글 입력 폴더"
        write_eml(
            input_root / "메일 001.eml",
            subject="Batch run",
            email_rows=[["교육명", "AI Basic"]],
            docx_rows=[["과정 평가", "만족"]],
        )
        working_dir = self.root / "unrelated cwd"
        working_dir.mkdir()
        command = f'cmd.exe /d /c call "{RUN_BAT}" "{input_root}"'
        result = subprocess.run(
            command,
            cwd=working_dir,
            input="\n",
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=30,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stdout)
        self.assertIn(str(EMBEDDED_PYTHON), result.stdout)
        self.assertTrue((input_root / app.OUTPUT_FILENAME).is_file())

    @unittest.skipUnless(EMBEDDED_PYTHON.is_file(), "embedded Python is not configured")
    def test_run_bat_accepts_console_folder_input(self) -> None:
        input_root = self.root / "console input"
        write_eml(
            input_root / "001.eml",
            subject="Console run",
            email_rows=[["교육명", "AI Basic"]],
            docx_rows=[["과정 평가", "만족"]],
        )
        working_dir = self.root / "another cwd"
        working_dir.mkdir()
        command = f'cmd.exe /d /c call "{RUN_BAT}"'
        result = subprocess.run(
            command,
            cwd=working_dir,
            input=f'"{input_root}"\n\n',
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            timeout=30,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stdout)
        self.assertTrue((input_root / app.OUTPUT_FILENAME).is_file())


if __name__ == "__main__":
    unittest.main()
