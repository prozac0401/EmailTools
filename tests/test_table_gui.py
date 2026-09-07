from __future__ import annotations

import base64
import http.client
import importlib.util
import json
import subprocess
import sys
import tempfile
import threading
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("emailtools_gui", ROOT / "eml_table_to_excel" / "gui.py")
gui = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = gui
SPEC.loader.exec_module(gui)
from openpyxl import load_workbook


class PreviewTests(unittest.TestCase):
    def setUp(self):
        self.workspace = gui.Workspace()

    def tearDown(self):
        self.workspace.close()

    def load(self, request):
        self.workspace.load(request)
        self.workspace.worker.join(timeout=10)
        self.assertFalse(self.workspace.worker.is_alive())
        return self.workspace.snapshot()

    def test_scan_and_column_selection_do_not_write_or_mutate_sources(self):
        source = self.workspace.root / "source"
        source.mkdir()
        gui.create_demo(source)
        before = {p.name: p.read_bytes() for p in source.iterdir()}
        state = self.load({"source": str(source)})
        self.assertEqual(state["phase"], "ready")
        self.assertEqual((state["total"], state["success"], state["warnings"]), (3, 2, 1))
        self.assertEqual({p.name: p.read_bytes() for p in source.iterdir()}, before)
        specs = [
            {"source": "교육명", "name": "교육명", "enabled": True},
            {"source": "강사", "name": "강사", "enabled": False},
            {"source": None, "name": "검토 상태", "value": "검토 대기", "enabled": True},
            {"source": "참석 인원", "name": "참석 인원", "enabled": True},
        ]
        request = {"dataset": state["dataset"], "columns": specs}
        preview = self.workspace.preview(request)
        self.assertEqual(preview["columns"], ["교육명", "검토 상태", "참석 인원"])
        self.assertEqual(preview["rows"][0], {"교육명": "AI 업무 활용", "검토 상태": "검토 대기", "참석 인원": ""})
        self.assertEqual(preview["rows"][1]["참석 인원"], "24")
        self.assertIn("강사", self.workspace.scan.records[0])
        result = self.workspace.export(request)
        exported = self.workspace.exports[result["download"].split("/")[1]] / "tables" / gui.core.OUTPUT_FILENAME
        workbook = load_workbook(exported)
        try:
            rows = list(workbook.active.values)
            self.assertEqual(list(rows[0]), preview["columns"])
            self.assertEqual([list(map(lambda v: v or "", row)) for row in rows[1:]],
                             [list(row.values()) for row in preview["rows"]])
        finally:
            workbook.close()
        self.assertEqual({p.name: p.read_bytes() for p in source.iterdir()}, before)
        # Restoring an excluded column recovers original data, and reordering is exact.
        specs[1]["enabled"] = True
        specs.reverse()
        restored = self.workspace.preview(request)
        self.assertEqual(restored["columns"], ["참석 인원", "검토 상태", "강사", "교육명"])
        self.assertEqual(restored["rows"][0]["강사"], "김민지")

    def test_preview_matches_export_for_literal_entities_formula_and_truncation(self):
        state = self.load({"demo": True})
        request = {"dataset": state["dataset"], "columns": [
            {"source": None, "name": "엔티티", "value": "&amp;lt;tag&amp;gt;", "enabled": True},
            {"source": None, "name": "수식 모양", "value": '=HYPERLINK("https://example.invalid")', "enabled": True},
            {"source": None, "name": "긴 값", "value": "X" * 40000 + "\x01", "enabled": True},
        ]}
        preview = self.workspace.preview(request)
        result = self.workspace.export(request)
        folder = self.workspace.exports[result["download"].split("/")[1]]
        workbook = load_workbook(folder / "tables" / gui.core.OUTPUT_FILENAME)
        try:
            self.assertEqual([cell.value for cell in workbook.active[2]], list(preview["rows"][0].values()))
            self.assertEqual(workbook.active["B2"].data_type, "s")
            self.assertEqual(len(workbook.active["C2"].value), 32767)
        finally:
            workbook.close()
        self.assertIn("truncated", (folder / "processing.log").read_text(encoding="utf-8-sig"))

    def test_invalid_columns_and_stale_dataset_rejected(self):
        state = self.load({"demo": True})
        valid = {"source": "교육명", "name": "교육명", "enabled": True}
        invalids = [[], [{**valid, "enabled": False}], [valid, valid],
                    [{**valid, "source": "missing"}], [{**valid, "name": " "}],
                    [{**valid, "name": "multi\nline"}], [{**valid, "name": "A" * 121}],
                    [{**valid, "enabled": "yes"}], [valid, {"name": "other", "source": "교육명", "enabled": True}]]
        for columns in invalids:
            with self.subTest(columns=columns):
                with self.assertRaises(ValueError):
                    self.workspace.preview({"dataset": state["dataset"], "columns": columns})
        self.load({"demo": True})
        with self.assertRaises(ValueError):
            self.workspace.export({"dataset": state["dataset"], "columns": [valid]})

    def test_upload_nested_korean_paths_and_failed_reload_clear_old_data(self):
        fixture = self.workspace.root / "fixtures"
        fixture.mkdir()
        gui.create_demo(fixture)
        payload = next(fixture.glob("*.eml")).read_bytes()
        state = self.load({"files": [{"name": "한글 폴더/첨부.eml", "data": base64.b64encode(payload).decode()}]})
        self.assertEqual(state["phase"], "ready")
        self.assertIn("한글 폴더", self.workspace.scan.records[0]["source_eml"])
        state = self.load({"source": str(fixture / "missing")})
        self.assertEqual(state["phase"], "error")
        self.assertNotIn("columns", state)
        self.assertIsNone(self.workspace.scan)
        self.assertEqual(state["dataset"], "")

    def test_upload_rejects_traversal_duplicates_and_non_eml(self):
        for names in [["../escape.eml"], ["C:/escape.eml"], ["/escape.eml"], ["run.exe"], ["mail.eml", "MAIL.eml"]]:
            with self.subTest(names=names), tempfile.TemporaryDirectory() as tmp:
                with self.assertRaises(ValueError):
                    gui.save_uploads(Path(tmp), [{"name": n, "data": "YWJj"} for n in names])

    def test_pagination_does_not_truncate_export(self):
        state = self.load({"demo": True})
        self.workspace.scan.records = [{"교육명": str(i)} for i in range(61)]
        request = {"dataset": state["dataset"], "columns": [{"source": "교육명", "name": "교육명", "enabled": True}], "page": 2}
        data = self.workspace.preview(request)
        self.assertEqual((len(data["rows"]), data["total"]), (11, 61))
        self.assertEqual(data["rows"][0]["교육명"], "50")
        result = self.workspace.export(request)
        self.assertEqual(result["total"], 61)
        folder = self.workspace.exports[result["download"].split("/")[1]]
        workbook = load_workbook(folder / "tables" / gui.core.OUTPUT_FILENAME)
        try:
            self.assertEqual(workbook.active.max_row, 62)
        finally:
            workbook.close()


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.workspace = gui.Workspace()
        self.server = gui.LocalServer(self.workspace)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.server.shutdown()
        self.thread.join()
        self.server.server_close()
        self.workspace.close()

    def request(self, method, route, body=None, headers=None):
        conn = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=10)
        try:
            conn.request(method, self.server.prefix + route, json.dumps(body) if body is not None else None, headers or {})
            response = conn.getresponse()
            return response.status, dict(response.headers), response.read()
        finally:
            conn.close()

    def test_http_preview_export_download_and_request_isolation(self):
        status, headers, body = self.request("GET", "")
        self.assertEqual(status, 200)
        self.assertIn("메일 불러오기", body.decode())
        self.assertIn("frame-ancestors 'none'", headers["Content-Security-Policy"])
        self.assertEqual(self.request("POST", "api/load", {"demo": True})[0], 403)
        self.assertEqual(self.request("GET", "api/state", headers={"Host": "evil.invalid"})[0], 403)
        self.assertEqual(self.request("GET", "api/state", headers={"Origin": "https://evil.invalid"})[0], 403)
        auth = {"X-EmailTools-Token": self.server.token, "Content-Type": "application/json"}
        self.assertEqual(self.request("POST", "api/load", {"demo": True}, auth)[0], 202)
        self.workspace.worker.join(timeout=10)
        state = json.loads(self.request("GET", "api/state")[2])
        request = {"dataset": state["dataset"], "columns": [{"source": "교육명", "name": "교육명", "enabled": True}]}
        self.assertEqual(self.request("POST", "api/preview", request, auth)[0], 200)
        result = json.loads(self.request("POST", "api/export", request, auth)[2])
        status, headers, body = self.request("GET", result["download"])
        self.assertEqual(status, 200)
        self.assertTrue(body.startswith(b"PK"))
        self.assertIn("attachment", headers["Content-Disposition"])
        self.assertEqual(self.request("GET", "../gui.py")[0], 404)

    @unittest.skipUnless(sys.platform == "win32", "Windows BAT launcher")
    def test_bat_starts_gui_from_other_cwd_with_unicode_source_and_shuts_down(self):
        with tempfile.TemporaryDirectory(prefix="한글 입력 ") as tmp:
            gui.create_demo(Path(tmp))
            command = f'cmd.exe /d /c call "{ROOT / "run.bat"}" --no-browser "{tmp}"'
            process = subprocess.Popen(command, cwd=tmp, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                       text=True, encoding="utf-8", errors="replace")
            try:
                # Read in a helper so a launcher failure cannot hang the test runner.
                from queue import Queue
                lines = Queue()
                def read_lines():
                    for line in process.stdout:
                        lines.put(line)
                reader = threading.Thread(target=read_lines, daemon=True)
                reader.start()
                output = ""
                for _ in range(8):
                    line = lines.get(timeout=15)
                    output += line
                    if "http://127.0.0.1:" in line:
                        url = line[line.index("http://"):].strip()
                        break
                else:
                    self.fail(output)
                from urllib.parse import urlsplit
                address = urlsplit(url)
                conn = http.client.HTTPConnection(address.hostname, address.port, timeout=10)
                conn.request("GET", address.path + "api/state")
                state = json.loads(conn.getresponse().read())
                self.assertEqual(state["initial_source"], tmp)
                conn.close()
                conn = http.client.HTTPConnection(address.hostname, address.port, timeout=10)
                conn.request("POST", address.path + "api/shutdown", "{}",
                             {"X-EmailTools-Token": address.path.strip("/"), "Content-Type": "application/json"})
                self.assertEqual(conn.getresponse().status, 200)
                conn.close()
                process.wait(timeout=10)
                reader.join(timeout=5)
                self.assertEqual(process.returncode, 0, output)
            finally:
                if process.poll() is None:
                    subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True)
                    process.wait(timeout=10)
                process.stdout.close()


if __name__ == "__main__":
    unittest.main()
