from __future__ import annotations

import html
import importlib.util
import logging
import math
import os
import re
import sys
import tempfile
import traceback
import zipfile
from dataclasses import dataclass
from email import policy
from email.message import Message
from email.parser import BytesParser
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Callable, Iterable
import xml.etree.ElementTree as ET


APP_DIR = Path(__file__).resolve().parent
REPO_ROOT = APP_DIR.parent
VENDOR_DIR = APP_DIR / "vendor"
EML_SUPPORT_PATH = REPO_ROOT / "eml_attachment_tool" / "app.py"

if VENDOR_DIR.is_dir():
    sys.path.insert(0, str(VENDOR_DIR))

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover - exercised by deployment checks
    raise RuntimeError(
        "bundled openpyxl is missing; restore eml_table_to_excel/vendor"
    ) from exc


APP_NAME = "EML Table to Excel"
APP_VERSION = "1.0.0"
OUTPUT_FILENAME = "EML_Table_Result.xlsx"
ERROR_LOG_FILENAME = "EML_Table_Result_errors.log"
WORKSHEET_NAME = "Result"

SYSTEM_COLUMNS = [
    "source_eml",
    "_status",
    "_error",
    "subject",
    "from",
    "to",
    "date",
    "docx_filename",
]

# Intentionally empty for v1.  Real aliases can be added later without changing
# the parsing pipeline after they have been confirmed against actual data.
COLUMN_ALIASES: dict[str, str] = {}

DUPLICATE_VALUE_SEPARATOR = " | "
MAX_EXCEL_CELL_CHARS = 32767
MAX_DOCX_XML_BYTES = 32 * 1024 * 1024
ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
KEY_EDGE_CHARS = " \t\r\n:：;；|·•※*#"
DOCX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
HEADER_KEYS = {"항목", "항목명", "구분", "key", "field", "item"}
HEADER_VALUES = {"값", "내용", "value", "description"}


def load_eml_support():
    module_name = "emailtools_eml_attachment_support"
    if module_name in sys.modules:
        return sys.modules[module_name]
    if not EML_SUPPORT_PATH.is_file():
        raise FileNotFoundError(f"기존 EML 도구를 찾을 수 없습니다: {EML_SUPPORT_PATH}")
    spec = importlib.util.spec_from_file_location(module_name, EML_SUPPORT_PATH)
    if spec is None or spec.loader is None:
        raise ImportError(f"기존 EML 도구를 불러올 수 없습니다: {EML_SUPPORT_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


EML_SUPPORT = load_eml_support()


def normalize_key(text: object) -> str:
    value = html.unescape(str(text or "")).replace("\xa0", " ")
    value = re.sub(r"[\r\n\t]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip(KEY_EDGE_CHARS)
    value = re.sub(r"\s+", " ", value).strip()
    return COLUMN_ALIASES.get(value, value)


def normalize_value(text: object) -> str:
    value = html.unescape(str(text or "")).replace("\xa0", " ")
    value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    lines = [re.sub(r"[ \f\v]+", " ", line).strip() for line in value.split("\n")]
    compact: list[str] = []
    for line in lines:
        if line:
            compact.append(line)
        elif compact and compact[-1] != "":
            compact.append("")
    while compact and compact[-1] == "":
        compact.pop()
    return "\n".join(compact)


def merge_key_value(target: dict[str, str], key: object, value: object) -> None:
    normalized_key = normalize_key(key)
    if not normalized_key:
        return
    normalized_value = normalize_value(value)
    if normalized_key not in target:
        target[normalized_key] = normalized_value
        return
    current = target[normalized_key]
    if not normalized_value or normalized_value == current:
        return
    if not current:
        target[normalized_key] = normalized_value
        return
    existing_values = [part.strip() for part in current.split(DUPLICATE_VALUE_SEPARATOR)]
    if normalized_value not in existing_values:
        target[normalized_key] = current + DUPLICATE_VALUE_SEPARATOR + normalized_value


def _is_header_pair(key: str, value: str) -> bool:
    return key.casefold() in HEADER_KEYS and value.casefold() in HEADER_VALUES


def table_to_key_value_dict(table: Iterable[Iterable[object]]) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in table:
        cells = [normalize_value(cell) for cell in row]
        for index in range(0, len(cells) - 1, 2):
            key = normalize_key(cells[index])
            value = normalize_value(cells[index + 1])
            if not key or _is_header_pair(key, value):
                continue
            merge_key_value(result, key, value)
    return result


class HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table_depth = 0
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs) -> None:
        del attrs
        tag = tag.lower()
        if tag == "table":
            self._table_depth += 1
            if self._table_depth == 1:
                self._table = []
            return
        if self._table_depth != 1:
            return
        if tag == "tr":
            self._finish_row()
            self._row = []
        elif tag in {"td", "th"} and self._row is not None:
            self._finish_cell()
            self._cell_parts = []
        elif tag == "br" and self._cell_parts is not None:
            self._cell_parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag == "table":
            if self._table_depth == 1:
                self._finish_row()
                if self._table:
                    self.tables.append(self._table)
                self._table = None
            self._table_depth = max(0, self._table_depth - 1)
            return
        if self._table_depth != 1:
            return
        if tag in {"td", "th"}:
            self._finish_cell()
        elif tag == "tr":
            self._finish_row()
        elif tag in {"p", "div", "li"} and self._cell_parts is not None:
            self._cell_parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._cell_parts is not None:
            self._cell_parts.append(data)

    def close(self) -> None:
        super().close()
        if self._table_depth:
            self._finish_row()
            if self._table:
                self.tables.append(self._table)
            self._table = None
            self._table_depth = 0

    def _finish_cell(self) -> None:
        if self._cell_parts is None:
            return
        if self._row is not None:
            self._row.append(normalize_value("".join(self._cell_parts)))
        self._cell_parts = None

    def _finish_row(self) -> None:
        self._finish_cell()
        if self._row is not None and self._table is not None:
            if any(normalize_value(cell) for cell in self._row):
                self._table.append(self._row)
        self._row = None


def parse_html_tables(html_text: str) -> list[list[list[str]]]:
    parser = HTMLTableParser()
    parser.feed(html_text)
    parser.close()
    return parser.tables


def extract_email_table_values(msg: Message) -> tuple[dict[str, str], int]:
    values: dict[str, str] = {}
    table_count = 0
    for part in msg.walk():
        if part.is_multipart():
            continue
        disposition = (part.get_content_disposition() or "").lower()
        if disposition == "attachment" or part.get_filename():
            continue
        if part.get_content_type().lower() != "text/html":
            continue
        html_text = EML_SUPPORT.payload_text(part)
        tables = parse_html_tables(html_text)
        table_count += len(tables)
        for table in tables:
            for key, value in table_to_key_value_dict(table).items():
                merge_key_value(values, key, value)
    return values, table_count


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_NAMESPACES = {"w": WORD_NS}


def _docx_cell_text(cell: ET.Element) -> str:
    paragraphs: list[str] = []
    for paragraph in cell.findall(".//w:p", WORD_NAMESPACES):
        parts: list[str] = []
        for element in paragraph.iter():
            local = element.tag.rsplit("}", 1)[-1]
            if local == "t" and element.text:
                parts.append(element.text)
            elif local == "tab":
                parts.append("\t")
            elif local in {"br", "cr"}:
                parts.append("\n")
        paragraph_text = normalize_value("".join(parts))
        if paragraph_text:
            paragraphs.append(paragraph_text)
    return normalize_value("\n".join(paragraphs))


def parse_docx_tables(docx_bytes: bytes) -> list[list[list[str]]]:
    with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
        info = archive.getinfo("word/document.xml")
        if info.file_size > MAX_DOCX_XML_BYTES:
            raise ValueError(
                f"word/document.xml is too large ({info.file_size} bytes; "
                f"limit {MAX_DOCX_XML_BYTES})"
            )
        root = ET.fromstring(archive.read(info))

    tables: list[list[list[str]]] = []
    for table_element in root.iter(f"{{{WORD_NS}}}tbl"):
        table: list[list[str]] = []
        for row_element in table_element.findall("./w:tr", WORD_NAMESPACES):
            row = [
                _docx_cell_text(cell)
                for cell in row_element.findall("./w:tc", WORD_NAMESPACES)
            ]
            if any(row):
                table.append(row)
        if table:
            tables.append(table)
    return tables


def header_value(msg: Message, name: str) -> str:
    return DUPLICATE_VALUE_SEPARATOR.join(
        normalize_value(str(value)) for value in msg.get_all(name, []) if str(value)
    )


def _record_data_key(key: str) -> str:
    if key in SYSTEM_COLUMNS:
        return f"item_{key}"
    return key


def _relative_source_name(eml_path: Path, source_root: Path) -> str:
    try:
        return str(eml_path.relative_to(source_root))
    except ValueError:
        return eml_path.name


@dataclass
class ProcessResult:
    record: dict[str, str]
    status: str
    errors: list[str]


def process_eml_record(
    eml_path: Path,
    source_root: Path,
    logger: logging.Logger,
) -> ProcessResult:
    source_name = _relative_source_name(eml_path, source_root)
    base_record = {column: "" for column in SYSTEM_COLUMNS}
    base_record["source_eml"] = source_name

    try:
        msg = BytesParser(policy=policy.default).parsebytes(eml_path.read_bytes())
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"
        base_record["_status"] = "PARSE ERROR"
        base_record["_error"] = error
        logger.exception("PARSE ERROR: %s", eml_path)
        return ProcessResult(base_record, "PARSE ERROR", [error])

    base_record.update(
        {
            "subject": header_value(msg, "Subject"),
            "from": header_value(msg, "From"),
            "to": header_value(msg, "To"),
            "date": header_value(msg, "Date"),
        }
    )

    warnings: list[str] = []
    errors: list[str] = []
    dynamic_values: dict[str, str] = {}

    try:
        email_values, email_table_count = extract_email_table_values(msg)
        if not email_table_count or not email_values:
            warnings.append("NO EMAIL TABLE")
        for key, value in email_values.items():
            merge_key_value(dynamic_values, _record_data_key(key), value)
    except Exception as exc:
        message = f"email table: {type(exc).__name__}: {exc}"
        errors.append(message)
        warnings.append("EMAIL TABLE ERROR")
        logger.exception("EMAIL TABLE ERROR: %s", eml_path)

    docx_names: list[str] = []
    docx_table_count = 0
    docx_parse_errors = 0
    for index, part in enumerate(EML_SUPPORT.attachment_parts(msg), 1):
        filename = part.get_filename() or ""
        content_type = part.get_content_type().lower()
        if Path(filename).suffix.lower() != ".docx" and content_type not in DOCX_CONTENT_TYPES:
            continue
        display_name = EML_SUPPORT.safe_name(filename, f"attachment_{index:03d}.docx")
        docx_names.append(display_name)
        try:
            payload = EML_SUPPORT.attachment_bytes(part)
            tables = parse_docx_tables(payload)
            docx_table_count += len(tables)
            for table in tables:
                for key, value in table_to_key_value_dict(table).items():
                    merge_key_value(dynamic_values, _record_data_key(key), value)
        except Exception as exc:
            docx_parse_errors += 1
            message = f"{display_name}: {type(exc).__name__}: {exc}"
            errors.append(message)
            logger.exception("DOCX PARSE ERROR: %s / %s", eml_path, display_name)

    for name in docx_names:
        merge_key_value(base_record, "docx_filename", name)

    if not docx_names:
        warnings.append("NO DOCX")
    elif docx_parse_errors:
        warnings.append("DOCX PARSE ERROR")
    elif not docx_table_count:
        warnings.append("NO DOCX TABLE")

    status = DUPLICATE_VALUE_SEPARATOR.join(dict.fromkeys(warnings)) or "OK"
    base_record["_status"] = status
    base_record["_error"] = "; ".join(errors)
    base_record.update(dynamic_values)

    if status != "OK":
        logger.warning("%s: %s%s", source_name, status, f"; {base_record['_error']}" if errors else "")
    return ProcessResult(base_record, status, errors)


def collect_columns(records: Iterable[dict[str, str]]) -> list[str]:
    columns = list(SYSTEM_COLUMNS)
    seen = set(columns)
    for record in records:
        for key in record:
            if key not in seen:
                columns.append(key)
                seen.add(key)
    return columns


def safe_excel_value(
    value: object,
    logger: logging.Logger,
    context: str,
) -> str:
    text = normalize_value(value)
    cleaned = ILLEGAL_XML_CHARS.sub(" ", text)
    if cleaned != text:
        logger.warning("Excel illegal control characters replaced: %s", context)
    if len(cleaned) > MAX_EXCEL_CELL_CHARS:
        logger.warning(
            "Excel cell truncated from %d to %d characters: %s",
            len(cleaned),
            MAX_EXCEL_CELL_CHARS,
            context,
        )
        return cleaned[: MAX_EXCEL_CELL_CHARS - 3] + "..."
    return cleaned


def _display_width(value: object) -> int:
    lines = str(value or "").splitlines() or [""]
    return max(sum(2 if ord(char) > 127 else 1 for char in line) for line in lines)


def _force_literal_text(cell) -> None:
    # openpyxl interprets strings beginning with '=' as formulas.  Every value
    # in this workbook comes from an untrusted email and must remain inert data.
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"


def write_excel(
    records: list[dict[str, str]],
    columns: list[str],
    output_path: Path,
    logger: logging.Logger,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = WORKSHEET_NAME
    worksheet.sheet_view.showGridLines = False

    worksheet.append(columns)
    for row_number, record in enumerate(records, 2):
        worksheet.append(
            [
                safe_excel_value(
                    record.get(column, ""),
                    logger,
                    f"{record.get('source_eml', '')} / {column}",
                )
                for column in columns
            ]
        )

    last_column = get_column_letter(len(columns))
    last_row = max(1, len(records) + 1)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{last_column}{last_row}"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        _force_literal_text(cell)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 28

    for row in worksheet.iter_rows(min_row=2, max_row=last_row):
        for cell in row:
            _force_literal_text(cell)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    preferred_widths = {
        "source_eml": 28,
        "_status": 24,
        "_error": 40,
        "subject": 32,
        "from": 26,
        "to": 26,
        "date": 24,
        "docx_filename": 28,
    }
    column_widths: dict[int, float] = {}
    for column_index, column_name in enumerate(columns, 1):
        measured = max(
            _display_width(worksheet.cell(row=row_index, column=column_index).value)
            for row_index in range(1, last_row + 1)
        )
        width = max(10, preferred_widths.get(column_name, 10), measured + 2)
        width = min(50, width)
        column_widths[column_index] = width
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(2, last_row + 1):
        estimated_lines = 1
        for column_index in range(1, len(columns) + 1):
            value = str(worksheet.cell(row=row_index, column=column_index).value or "")
            width = max(1, column_widths[column_index])
            wrapped_lines = sum(
                max(1, math.ceil(_display_width(line) / width))
                for line in (value.splitlines() or [""])
            )
            estimated_lines = max(estimated_lines, wrapped_lines)
        worksheet.row_dimensions[row_index].height = min(90, max(18, estimated_lines * 15))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}_",
        suffix=".tmp.xlsx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
        workbook.close()


@dataclass
class BatchResult:
    output_path: Path
    log_path: Path
    total: int
    success: int
    warnings: int
    failures: int
    columns: list[str]
    records: list[dict[str, str]]


def find_eml_files(source_root: Path) -> list[Path]:
    return sorted(
        (
            path
            for path in source_root.rglob("*")
            if path.is_file() and path.suffix.lower() == ".eml"
        ),
        key=lambda path: str(path.relative_to(source_root)).casefold(),
    )


def configure_logger(log_path: Path) -> logging.Logger:
    logger = logging.getLogger("emailtools.eml_table_to_excel")
    logger.handlers.clear()
    logger.propagate = False
    logger.setLevel(logging.INFO)
    handler = logging.FileHandler(log_path, mode="w", encoding="utf-8-sig")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(handler)
    return logger


def close_logger(logger: logging.Logger) -> None:
    for handler in list(logger.handlers):
        handler.close()
        logger.removeHandler(handler)


def run_batch(
    source_root: Path,
    progress: Callable[[str], None] | None = None,
) -> BatchResult:
    source_root = source_root.expanduser().resolve()
    if not source_root.is_dir():
        raise NotADirectoryError(f"EML 폴더를 찾을 수 없습니다: {source_root}")

    eml_files = find_eml_files(source_root)
    if not eml_files:
        raise FileNotFoundError(f"EML 파일을 찾지 못했습니다: {source_root}")

    output_path = source_root / OUTPUT_FILENAME
    log_path = source_root / ERROR_LOG_FILENAME
    logger = configure_logger(log_path)
    records: list[dict[str, str]] = []
    success = 0
    warning_count = 0
    failures = 0
    try:
        logger.info("Start: %s (%d EML files)", source_root, len(eml_files))
        for index, eml_path in enumerate(eml_files, 1):
            try:
                result = process_eml_record(eml_path, source_root, logger)
            except Exception as exc:  # Defensive batch boundary.
                error = f"{type(exc).__name__}: {exc}"
                logger.exception("UNEXPECTED ERROR: %s", eml_path)
                record = {column: "" for column in SYSTEM_COLUMNS}
                record["source_eml"] = _relative_source_name(eml_path, source_root)
                record["_status"] = "PARSE ERROR"
                record["_error"] = error
                result = ProcessResult(record, "PARSE ERROR", [error])

            records.append(result.record)
            if result.status == "OK":
                success += 1
            elif result.status == "PARSE ERROR":
                failures += 1
            else:
                warning_count += 1
            if progress:
                progress(f"[{index}/{len(eml_files)}] {eml_path.name} ... {result.status}")

        columns = collect_columns(records)
        write_excel(records, columns, output_path, logger)
        logger.info(
            "Complete: total=%d success=%d warnings=%d failures=%d output=%s",
            len(eml_files),
            success,
            warning_count,
            failures,
            output_path,
        )
    finally:
        close_logger(logger)

    return BatchResult(
        output_path=output_path,
        log_path=log_path,
        total=len(eml_files),
        success=success,
        warnings=warning_count,
        failures=failures,
        columns=columns,
        records=records,
    )


def choose_source_folder() -> Path:
    print()
    print("EML 파일이 있는 폴더를 입력하세요.")
    print("예: D:\\Mail\\Training")
    raw = input("> ").strip().strip('"')
    if not raw:
        raise ValueError("폴더 경로가 입력되지 않았습니다.")
    return Path(raw)


def main() -> int:
    print("=" * 72)
    print(f"{APP_NAME} v{APP_VERSION}")
    print("Shared Windows Embedded Python / bundled openpyxl")
    print("=" * 72)
    try:
        source_root = Path(sys.argv[1].strip('"')) if len(sys.argv) > 1 else choose_source_folder()
        result = run_batch(source_root, print)
        print()
        print("=" * 72)
        print("처리 완료")
        print(f"전체: {result.total}")
        print(f"성공: {result.success}")
        print(f"경고: {result.warnings}")
        print(f"실패: {result.failures}")
        print()
        print(f"결과: {result.output_path}")
        print(f"로그: {result.log_path}")
        print("=" * 72)
        return 0
    except Exception as exc:
        print(f"\n[오류] {type(exc).__name__}: {exc}")
        print(f"상세 로그를 생성할 수 있는 입력 폴더인지 확인해 주세요.")
        if os.environ.get("EMAILTOOLS_DEBUG") == "1":
            traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
